"""Make text boxes actually fit the text they hold.

python-pptx cannot lay text out, so a box's declared height is whatever
the caller guessed when creating it. These decks build a box with
add_text(..., height=0.5) and then append paragraphs to it, so the
declared height is almost always wrong: on Update 12 one box declared
0.5in while its text needed 3.1in.

PowerPoint still draws the overflow, but the shape's handles sit where
the declared height says, so editing the deck by hand means resizing
every box first. That is the problem this fixes.

fit_text_boxes() measures each paragraph with real font metrics, wraps
it at the box's usable width, and rewrites the shape height to what the
text needs. It also sets spAutoFit so PowerPoint keeps the height
correct if the text is edited later.

Arial is the metric proxy for Calibri. Calibri is the narrower face, so
measuring with Arial errs toward a slightly taller box, never a short one.
"""
from pptx.util import Pt, Emu
from pptx.enum.text import MSO_AUTO_SIZE

try:
    from PIL import ImageFont
    _PIL = True
except ImportError:
    _PIL = False

_REG = "/System/Library/Fonts/Supplemental/Arial.ttf"
_BOLD = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
_EMU_IN = 914400.0
_INSET_LR = 0.1      # PowerPoint default left/right inset, inches
_INSET_TB = 0.05     # default top/bottom inset
_LINE = 1.22         # line height as a multiple of font size
_RES = 4             # render scale for metric resolution

_cache = {}


def _font(size_pt, bold):
    key = (round(size_pt, 1), bool(bold))
    if key not in _cache:
        _cache[key] = ImageFont.truetype(_BOLD if bold else _REG, int(size_pt * _RES))
    return _cache[key]


def _width_pt(s, size_pt, bold):
    if not s:
        return 0.0
    # Non-Latin glyphs (emoji) fall back to .notdef in Arial and measure
    # narrow, so charge them a full em to stay conservative.
    extra = sum(size_pt for ch in s if ord(ch) > 0x2000)
    return _font(size_pt, bold).getlength(s) / _RES + extra


def _line_count(s, size_pt, bold, avail_pt):
    if not s.strip():
        return 1
    lines, cur = 1, ""
    for word in s.split():
        trial = word if not cur else cur + " " + word
        if _width_pt(trial, size_pt, bold) <= avail_pt or not cur:
            cur = trial
        else:
            lines += 1
            cur = word
    return lines


def measure(tf, width_in, default_pt=18):
    """Height in inches that this text frame's content needs."""
    ml = tf.margin_left / _EMU_IN if tf.margin_left is not None else _INSET_LR
    mr = tf.margin_right / _EMU_IN if tf.margin_right is not None else _INSET_LR
    mt = tf.margin_top / _EMU_IN if tf.margin_top is not None else _INSET_TB
    mb = tf.margin_bottom / _EMU_IN if tf.margin_bottom is not None else _INSET_TB
    avail_pt = (width_in - ml - mr) * 72
    total = 0.0
    for p in tf.paragraphs:
        s = "".join(r.text for r in p.runs) or p.text
        size = (p.font.size or Pt(default_pt)).pt
        bold = bool(p.font.bold)
        before = p.space_before.pt if p.space_before is not None else 0.0
        total += before + _line_count(s, size, bold, avail_pt) * size * _LINE
    return total / 72 + mt + mb


def fit_text_boxes(prs, grow_only=True, verbose=True):
    """Resize every text box so its declared height matches its content.

    grow_only leaves boxes that are already tall enough alone, so
    deliberate spacing is preserved and only overflow is corrected.
    """
    if not _PIL:
        print("  [pptx_fit] PIL unavailable, skipping fit pass")
        return 0
    fixed = 0
    for idx, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            tf = sh.text_frame
            if not tf.text.strip():
                continue
            need_in = measure(tf, sh.width / _EMU_IN)
            have_in = sh.height / _EMU_IN
            if need_in - have_in > 0.02 or not grow_only:
                if verbose:
                    print(f"  [pptx_fit] slide {idx}: {have_in:.2f}in -> {need_in:.2f}in  "
                          f"\"{tf.text.splitlines()[0][:44]}\"")
                sh.height = Emu(int(need_in * _EMU_IN))
                fixed += 1
            try:
                tf.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
            except Exception:
                pass
    return fixed
