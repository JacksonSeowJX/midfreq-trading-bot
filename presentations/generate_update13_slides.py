"""Generate FYP Update 13 Presentation — Settling Whether the Result Was Luck

Update 12 ended by promising to test whether the one passing result (cross-
sectional reversal on the S&P 100) was a real edge or the luckiest of many
combinations tried. This deck reports that test (White's Reality Check), the
finding that the result flips between pass and fail with a one-week shift in
the data window, the Hong Kong fee checked against a real trade, and the
live forward test of the strategy.

Reality Check figures are read from results/reality_check_2*.json so the
deck cannot drift from the study output.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path
import glob
import json
import sys

import pandas as pd

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BG = RGBColor(0xED, 0xE8, 0xE0)
BG_CARD = RGBColor(0xE0, 0xDB, 0xD3)
DARK = RGBColor(0x2D, 0x2D, 0x2D)
BRACKET = RGBColor(0x33, 0x33, 0x2D)
ACCENT_GREEN = RGBColor(0x2E, 0x7D, 0x32)
ACCENT_RED = RGBColor(0xC6, 0x28, 0x28)
DIM = RGBColor(0x6B, 0x6B, 0x63)
MID = RGBColor(0x55, 0x55, 0x50)

CHART_DIR = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent
RESULTS = Path(__file__).resolve().parent.parent / 'results'

_rc_file = sorted(glob.glob(str(RESULTS / 'reality_check_2*.json')))[-1]
RC = json.loads(Path(_rc_file).read_text())['by_config']
A, B = RC['A (9 windows)'], RC['B (15 windows)']

def set_slide_bg(slide, color=BG):
    f = slide.background.fill; f.solid(); f.fore_color.rgb = color


def add_text(slide, left, top, width, height, text, font_size=18, color=DARK, bold=False,
             alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text
    p.font.size = Pt(font_size); p.font.color.rgb = color
    p.font.bold = bold; p.font.name = font_name; p.alignment = alignment
    return tf


def add_para(tf, text, font_size=18, color=DARK, bold=False, alignment=PP_ALIGN.LEFT,
             font_name="Calibri", space_before=Pt(6)):
    p = tf.add_paragraph(); p.text = text
    p.font.size = Pt(font_size); p.font.color.rgb = color
    p.font.bold = bold; p.font.name = font_name
    p.alignment = alignment; p.space_before = space_before
    return p


def add_bracket_tl(slide, left, top, size=1.2, thickness=0.12):
    v = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(thickness), Inches(size))
    v.fill.solid(); v.fill.fore_color.rgb = BRACKET; v.line.fill.background()
    h = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(size), Inches(thickness))
    h.fill.solid(); h.fill.fore_color.rgb = BRACKET; h.line.fill.background()


def add_bracket_br(slide, right, bottom, size=1.2, thickness=0.12):
    v = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(right - thickness), Inches(bottom - size), Inches(thickness), Inches(size))
    v.fill.solid(); v.fill.fore_color.rgb = BRACKET; v.line.fill.background()
    h = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(right - size), Inches(bottom - thickness), Inches(size), Inches(thickness))
    h.fill.solid(); h.fill.fore_color.rgb = BRACKET; h.line.fill.background()


def add_divider(slide, left, top, width):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Pt(2))
    s.fill.solid(); s.fill.fore_color.rgb = BRACKET; s.line.fill.background()


def add_card(slide, left, top, width, height, fill_color=BG_CARD):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    s.fill.solid(); s.fill.fore_color.rgb = fill_color
    s.line.color.rgb = RGBColor(0xCC, 0xC7, 0xBF); s.line.width = Pt(1)
    return s


def picture(slide, name, left, top, width):
    img = CHART_DIR / name
    if img.exists():
        slide.shapes.add_picture(str(img), Inches(left), Inches(top), width=Inches(width))




def title(slide, text, size=28):
    add_bracket_tl(slide, 0.5, 0.3, size=0.8, thickness=0.1)
    add_text(slide, 0.8, 0.5, 12, 0.7, text, font_size=size, color=DARK, bold=True)
    add_divider(slide, 0.8, 1.15, 4)


def caption_card(slide, top, headline, body, height=1.1):
    add_card(slide, 0.8, top, 11.7, height)
    tf = add_text(slide, 1.1, top + 0.1, 11.2, 0.5, headline, font_size=14, color=DARK, bold=True)
    add_para(tf, body, font_size=13.5, color=MID, space_before=Pt(2))
    return tf


# ==================== SLIDE 1: TITLE ====================
slide = prs.slides.add_slide(prs.slide_layouts[6]); set_slide_bg(slide)
add_bracket_tl(slide, 1.5, 1.0, size=2.0, thickness=0.15)
add_bracket_br(slide, 11.8, 6.5, size=2.0, thickness=0.15)
add_text(slide, 2.5, 2.6, 8.5, 1.2, "FYP UPDATE 13", font_size=44, color=DARK, bold=True, alignment=PP_ALIGN.CENTER)
add_text(slide, 2.5, 3.8, 8.5, 0.8, "Was the Passing Result Luck?", font_size=20, color=MID, alignment=PP_ALIGN.CENTER)
add_text(slide, 2.5, 5.0, 8.5, 0.5, "Jackson Seow  •  FYP 2025/2026", font_size=16, color=DIM, alignment=PP_ALIGN.CENTER)


# ==================== SLIDE 2: WHERE UPDATE 12 LEFT IT ====================
slide = prs.slides.add_slide(prs.slide_layouts[6]); set_slide_bg(slide)
title(slide, "Where Update 12 Left It")
add_card(slide, 0.8, 1.6, 5.6, 4.4)
tf = add_text(slide, 1.1, 1.85, 5.0, 0.5, "What had passed", font_size=18, color=ACCENT_GREEN, bold=True)
for line in ("", "Cross-sectional reversal on the S&P 100:",
             "•  passed the 9 and 15 window standard",
             "•  survived the real broker fee of 0.005%",
             "•  PBO of 0.274, so the settings search",
             "    was not just picking noise"):
    add_para(tf, line, font_size=14, color=DARK)
add_card(slide, 6.9, 1.6, 5.6, 4.4, fill_color=RGBColor(0xE4, 0xDE, 0xD2))
tf = add_text(slide, 7.2, 1.85, 5.0, 0.5, "The question still open", font_size=18, color=ACCENT_RED, bold=True)
for line in ("", "Many strategy and universe combinations",
             "were tried and only 1 passed.", "",
             "Try enough combinations and one will pass",
             "by luck alone, so is this a real edge, or",
             "the lucky one?"):
    add_para(tf, line, font_size=14, color=DARK)
add_card(slide, 0.8, 6.25, 11.7, 0.75)
add_text(slide, 1.1, 6.37, 11.2, 0.5,
         "This update runs the test that answers it: White's Reality Check.", font_size=13.5, color=MID)


# ==================== SLIDE 3: HOW THE REALITY CHECK WORKS ====================
slide = prs.slides.add_slide(prs.slide_layouts[6]); set_slide_bg(slide)
title(slide, "Why the Number of Tries Matters")
picture(slide, 'chart_u13_cost_curve.png', 1.97, 1.4, 9.4)
caption_card(slide, 6.1,
             "The Reality Check: remove any edge from all 8 combinations, reshuffle 20,000 times, "
             "and count how often the luckiest one beats our winner.",
             "8 combinations = 2 cross-sectional strategies × 4 stock universes, all over the same 3 years. "
             "Counting every single-stock test instead would mean over 100 candidates and a worse score for our result.",
             height=1.2)


# ==================== SLIDE 4: THE RESULT ====================
slide = prs.slides.add_slide(prs.slide_layouts[6]); set_slide_bg(slide)
title(slide, "The Result: It Does Not Pass")
picture(slide, 'chart_u13_rc.png', 0.9, 1.4, 11.5)
caption_card(slide, 5.75,
             f"Luck matches or beats the winner {A['p_value']:.1%} and {B['p_value']:.1%} of the time. "
             f"The usual bar for \"real\" is under 5%.",
             f"Tested alone, the same winners score {A['naive_p_value']:.3f} and {B['naive_p_value']:.3f}. "
             "The gap between those numbers is the cost of admitting how many combinations were tried.",
             height=1.25)


# ==================== SLIDE 5: THE RESULT MOVES WITH THE DATA ====================
slide = prs.slides.add_slide(prs.slide_layouts[6]); set_slide_bg(slide)
title(slide, "The Result Also Changes With the Data Window")
picture(slide, 'chart_u13_drift.png', 1.6, 1.4, 10.1)
caption_card(slide, 6.05,
             "Moving the end of the 3-year window by one week, 21 to 28 August, turns a pass into a fail.",
             "The Hong Kong results, whose data did not change, came out exactly the same, "
             "so the change comes from the data, not the code.",
             height=1.15)


# ==================== SLIDE 6: HK FEE CHECKED ====================
slide = prs.slides.add_slide(prs.slide_layouts[6]); set_slide_bg(slide)
title(slide, "The Trading Fees, Checked Against Real Trades")
picture(slide, 'chart_u13_fee.png', 0.9, 1.4, 11.5)
tf = caption_card(slide, 5.75,
                  "A real paper trade of 100 Tencent shares was charged HK$77.73. "
                  "The published fee schedule predicts HK$77.84.",
                  "Because of the flat HK$15 platform fee, the percentage depends on trade size: "
                  "about 0.158% at the HK$100,000 our studies trade, against the 0.16% the code uses.",
                  height=1.6)
add_para(tf, "US: the first live orders (5 October) were charged US$0.99 each, 0.003% of a "
             "US$33,000 order, against the 0.005% the code uses.",
         font_size=13.5, color=MID, space_before=Pt(4))


# ==================== SLIDE 7: SUMMARY ====================
slide = prs.slides.add_slide(prs.slide_layouts[6]); set_slide_bg(slide)
add_bracket_tl(slide, 1.5, 0.7, size=2.0, thickness=0.15)
add_bracket_br(slide, 11.8, 6.8, size=2.0, thickness=0.15)
add_text(slide, 2.5, 1.0, 8.5, 0.7, "SUMMARY", font_size=32, color=DARK, bold=True, alignment=PP_ALIGN.CENTER)
tf = add_text(slide, 2.0, 2.0, 9.5, 0.5,
              f"❌   Reality Check: luck beats the winner {A['p_value']:.1%} and {B['p_value']:.1%} of the time, "
              "above the 5% bar", font_size=16, color=DARK)
add_para(tf, "", font_size=7)
add_para(tf, "❌   A one-week change in the data turns the pass into a fail", font_size=16, color=DARK)
add_para(tf, "", font_size=7)
add_para(tf, "✅   Fees confirmed by real trades: HK$77.73 vs HK$77.84 predicted, US$0.99 per US order", font_size=16, color=DARK)
add_para(tf, "", font_size=7)
add_para(tf, "➡️   No strategy in this project beat chance once tested properly", font_size=16, color=DARK, bold=True)
add_card(slide, 2.0, 4.6, 9.5, 2.05)
tf = add_text(slide, 2.3, 4.7, 9, 0.5, "🎯  Now trading live: S&P 100 reversal, US$100,000 paper account",
              font_size=17, color=ACCENT_GREEN, bold=True)
add_para(tf, "Live since 28 September, but week 1 placed no trades: the server's daily code update kept "
             "resetting the saved rebalance counter. Fixed, and the whole live path replayed against a "
             "simulated broker.", font_size=13.5, color=MID)
add_para(tf, "First rebalance on 5 October bought Qualcomm, General Motors and Netflix, about US$33,000 each, "
             "matching the broker share for share. Then one every 19 trading hours (just under 3 trading days): "
             "about 4 by 19 October.",
         font_size=13.5, color=MID, space_before=Pt(4))
add_text(slide, 2.5, 6.8, 8.5, 0.5, "Thank You", font_size=24, color=BRACKET, bold=True, alignment=PP_ALIGN.CENTER)


from pptx_fit import fit_text_boxes
print("fitting text boxes to their content:")
n = fit_text_boxes(prs)
print(f"  {n} box(es) resized")

output_path = "/Users/jacksonetherchainstake/FYP/Documents/FYP_Update_13.pptx"
prs.save(output_path)
print(f"Saved to {output_path}")
