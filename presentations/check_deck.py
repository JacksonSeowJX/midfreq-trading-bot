"""
Layout check for a generated deck: things python-pptx cannot see.

python-pptx places shapes where it is told and never lays anything out, so
a deck can build cleanly while text spills out of its card, a chart sits on
top of the caption card below it, or a shape hangs off the slide. Each of
these has happened in this project's decks and was only found by opening
them. Run after every build.

Usage:
    python3 presentations/check_deck.py path/to/deck.pptx
"""
import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

E = 914400.0


def box(sh):
    return sh.left / E, sh.top / E, sh.width / E, sh.height / E


def main(path):
    prs = Presentation(path)
    W, H = prs.slide_width / E, prs.slide_height / E
    problems = 0
    for n, s in enumerate(prs.slides, 1):
        cards = [box(sh) for sh in s.shapes
                 if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and box(sh)[2] > 3 and box(sh)[3] > 0.6]
        for sh in s.shapes:
            l, t, w, h = box(sh)
            label = (sh.text_frame.text.splitlines()[0][:45] if sh.has_text_frame and sh.text_frame.text.strip()
                     else str(sh.shape_type).split('.')[-1])
            if l + w > W + 0.02 or t + h > H + 0.02 or l < -0.02 or t < -0.02:
                print(f"  slide {n}: OFF SLIDE  {label}"); problems += 1
            if sh.has_text_frame and sh.text_frame.text.strip():
                host = [c for c in cards if c[0] - 0.05 <= l < c[0] + c[2] and c[1] - 0.05 <= t < c[1] + c[3]]
                if host and (t + h) - (host[0][1] + host[0][3]) > 0.03:
                    print(f"  slide {n}: text {((t+h)-(host[0][1]+host[0][3])):.2f}in past its card  {label}")
                    problems += 1
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                for c in cards:
                    if c[1] > t and t + h > c[1] + 0.02 and not (l + w <= c[0] or l >= c[0] + c[2]):
                        print(f"  slide {n}: picture overlaps the card below by {t + h - c[1]:.2f}in")
                        problems += 1
    print("layout OK" if not problems else f"{problems} layout problem(s)")
    return 1 if problems else 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1]))
